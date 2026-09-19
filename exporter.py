"""
Módulo de exportación de informes técnicos en formatos Excel (.xlsx) y PDF (.pdf)
con datos de entrada, tabla de resultados, curva granulométrica y parámetros geotécnicos.
"""

import os
import io
import tempfile
from typing import Optional
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

from models import ResultadoGranulometria
from plotting import generar_figura_curva_granulometrica


class NumberedCanvas(canvas.Canvas):
    """Lienzo personalizado para incluir encabezado y pie de página con numeración total."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#555555"))

        # Encabezado sutil
        self.drawString(36, 760, "Laboratorio de Mecánica de Suelos | Analizador Granulométrico Python")
        self.setStrokeColor(colors.HexColor("#CCCCCC"))
        self.setLineWidth(0.5)
        self.line(36, 754, 576, 754)

        # Pie de página
        self.line(36, 28, 576, 28)
        self.drawString(36, 18, f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - PySections Geotecnia")
        self.drawRightString(576, 18, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()


class ExportadorGranulometria:
    """Clase encargada de generar informes en Excel y PDF."""

    @staticmethod
    def exportar_excel(resultado: ResultadoGranulometria, ruta_archivo: str) -> str:
        """
        Exporta el análisis completo a un libro de Excel (.xlsx) estilizado
        con datos, tabla de tamices, parámetros y la curva granulométrica incrustada.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Análisis Granulométrico"
        ws.views.sheetView[0].showGridLines = True

        # Paleta de colores
        c_primario = "1B365D"    # Azul Marino
        c_secundario = "2C5E8A"  # Azul Medio
        c_claro = "EBF2F7"       # Fondo claro
        c_borde = "CCCCCC"

        f_titulo = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        f_subtitulo = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        f_seccion = Font(name="Calibri", size=11, bold=True, color=c_primario)
        f_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        f_bold = Font(name="Calibri", size=10, bold=True)
        f_normal = Font(name="Calibri", size=10)

        fill_titulo = PatternFill(start_color=c_primario, end_color=c_primario, fill_type="solid")
        fill_subtitulo = PatternFill(start_color=c_secundario, end_color=c_secundario, fill_type="solid")
        fill_claro = PatternFill(start_color=c_claro, end_color=c_claro, fill_type="solid")
        fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        borde_fino = Border(
            left=Side(style='thin', color=c_borde),
            right=Side(style='thin', color=c_borde),
            top=Side(style='thin', color=c_borde),
            bottom=Side(style='thin', color=c_borde)
        )
        borde_doble = Border(
            top=Side(style='thin', color=c_borde),
            bottom=Side(style='double', color=c_primario)
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # 1. TÍTULO PRINCIPAL
        ws.merge_cells("A1:G1")
        ws["A1"] = "INFORME DE ANÁLISIS GRANULOMÉTRICO POR TAMIZADO"
        ws["A1"].font = f_titulo
        ws["A1"].fill = fill_titulo
        ws["A1"].alignment = align_center
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A2:G2")
        ws["A2"] = "Norma Técnica ASTM C136 / ASTM D422 / ASTM D6913 - Clasificación SUCS ASTM D2487"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
        ws["A2"].fill = fill_subtitulo
        ws["A2"].alignment = align_center
        ws.row_dimensions[2].height = 20

        # 2. METADATOS DE LA MUESTRA
        row = 4
        ws.merge_cells(f"A{row}:G{row}")
        ws[f"A{row}"] = "DATOS DE IDENTIFICACIÓN Y MUESTRA"
        ws[f"A{row}"].font = f_seccion
        ws[f"A{row}"].fill = fill_claro
        row += 1

        metadatos = [
            ("Muestra ID:", resultado.datos_muestra.id_muestra, "Operador:", resultado.datos_muestra.operador),
            ("Proyecto:", resultado.datos_muestra.proyecto, "Fecha:", resultado.datos_muestra.fecha),
            ("Ubicación/Sondeo:", resultado.datos_muestra.ubicacion_sondeo, "Peso Total Inicial (g):", f"{resultado.peso_total_muestra_g:.2f}"),
            ("Curso / Cliente:", resultado.datos_muestra.cliente_o_curso, "Pérdida de Lavado / Pesaje:", f"{resultado.perdida_peso_g:.2f} g ({resultado.porcentaje_perdida:.2f}%)")
        ]

        for lab1, val1, lab2, val2 in metadatos:
            ws[f"A{row}"] = lab1
            ws[f"A{row}"].font = f_bold
            ws[f"B{row}"] = val1
            ws[f"B{row}"].font = f_normal

            ws[f"D{row}"] = lab2
            ws[f"D{row}"].font = f_bold
            ws[f"E{row}"] = val2
            ws[f"E{row}"].font = f_normal
            row += 1

        # 3. TABLA GRANULOMÉTRICA
        row += 1
        ws.merge_cells(f"A{row}:F{row}")
        ws[f"A{row}"] = "RESULTADOS DEL TAMIZADO"
        ws[f"A{row}"].font = f_seccion
        ws[f"A{row}"].fill = fill_claro
        row += 1

        headers = [
            "Tamiz",
            "Abertura (mm)",
            "Peso Retenido (g)",
            "% Retenido",
            "% Retenido Acum.",
            "% Que Pasa"
        ]

        ws.row_dimensions[row].height = 24
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = f_header
            cell.fill = fill_subtitulo
            cell.alignment = align_center
            cell.border = borde_fino

        fila_inicio_tabla = row + 1
        row += 1

        for idx, fila in enumerate(resultado.tabla):
            ws.row_dimensions[row].height = 19
            zebra = fill_zebra if idx % 2 == 1 else PatternFill(fill_type=None)

            c1 = ws.cell(row=row, column=1, value=fila.tamiz)
            c2 = ws.cell(row=row, column=2, value=fila.abertura_mm if fila.abertura_mm > 0 else "-")
            c3 = ws.cell(row=row, column=3, value=round(fila.peso_retenido_g, 2))
            c4 = ws.cell(row=row, column=4, value=round(fila.porcentaje_retenido, 2))
            c5 = ws.cell(row=row, column=5, value=round(fila.porcentaje_retenido_acumulado, 2))
            c6 = ws.cell(row=row, column=6, value=round(fila.porcentaje_pasa, 2))

            for c in [c1, c2, c3, c4, c5, c6]:
                c.font = f_normal
                c.border = borde_fino
                if zebra.fill_type:
                    c.fill = zebra

            c1.alignment = align_center
            c2.alignment = align_center
            c3.alignment = align_right
            c4.alignment = align_right
            c5.alignment = align_right
            c6.alignment = align_right

            row += 1

        # Fila de Totales
        ws.row_dimensions[row].height = 20
        c_tot_label = ws.cell(row=row, column=1, value="TOTAL")
        c_tot_label.font = f_bold
        c_tot_label.alignment = align_center
        c_tot_label.border = borde_doble

        c_tot_ab = ws.cell(row=row, column=2, value="-")
        c_tot_ab.alignment = align_center
        c_tot_ab.border = borde_doble

        c_tot_peso = ws.cell(row=row, column=3, value=round(resultado.peso_total_retenido_g, 2))
        c_tot_peso.font = f_bold
        c_tot_peso.alignment = align_right
        c_tot_peso.border = borde_doble

        c_tot_pct = ws.cell(row=row, column=4, value=round(sum(f.porcentaje_retenido for f in resultado.tabla), 2))
        c_tot_pct.font = f_bold
        c_tot_pct.alignment = align_right
        c_tot_pct.border = borde_doble

        for col in [5, 6]:
            c = ws.cell(row=row, column=col, value="-")
            c.alignment = align_center
            c.border = borde_doble

        row += 2

        # 4. PARÁMETROS GEOTÉCNICOS Y CLASIFICACIÓN
        ws.merge_cells(f"A{row}:F{row}")
        ws[f"A{row}"] = "PARÁMETROS GEOTÉCNICOS Y CLASIFICACIÓN DEL SUELO"
        ws[f"A{row}"].font = f_seccion
        ws[f"A{row}"].fill = fill_claro
        row += 1

        params_data = [
            ("Fracción Grava (%):", f"{resultado.porcentaje_grava:.2f}%", "Diámetro D10 (mm):", f"{resultado.d10:.3f}" if resultado.d10 else "N/D"),
            ("Fracción Arena (%):", f"{resultado.porcentaje_arena:.2f}%", "Diámetro D30 (mm):", f"{resultado.d30:.3f}" if resultado.d30 else "N/D"),
            ("  - Arena Gruesa (%):", f"{resultado.porcentaje_arena_gruesa:.2f}%", "Diámetro D60 (mm):", f"{resultado.d60:.3f}" if resultado.d60 else "N/D"),
            ("  - Arena Media (%):", f"{resultado.porcentaje_arena_media:.2f}%", "Coef. Uniformidad (Cu):", f"{resultado.cu:.2f}" if resultado.cu else "N/D"),
            ("  - Arena Fina (%):", f"{resultado.porcentaje_arena_fina:.2f}%", "Coef. Curvatura (Cc):", f"{resultado.cc:.2f}" if resultado.cc else "N/D"),
            ("Fracción Finos (%):", f"{resultado.porcentaje_finos:.2f}%", "Clasificación SUCS:", resultado.clasificacion_sucs),
            ("Descripción Suelo:", resultado.descripcion_suelo, "", "")
        ]

        for p_row in params_data:
            ws[f"A{row}"] = p_row[0]
            ws[f"A{row}"].font = f_bold
            ws[f"B{row}"] = p_row[1]
            ws[f"B{row}"].font = f_normal

            ws[f"D{row}"] = p_row[2]
            ws[f"D{row}"].font = f_bold
            ws[f"E{row}"] = p_row[3]
            ws[f"E{row}"].font = f_bold if "SUCS" in p_row[2] else f_normal

            if p_row[0] == "Descripción Suelo:":
                ws.merge_cells(f"B{row}:F{row}")
                ws[f"B{row}"].alignment = align_left

            row += 1

        # 5. GENERAR E INCRUSTAR CURVA GRANULOMÉTRICA
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_img:
            img_path = tmp_img.name

        try:
            fig = generar_figura_curva_granulometrica(resultado, figsize=(9.0, 5.0), dpi=130)
            fig.savefig(img_path, dpi=130, bbox_inches='tight')

            # Insertar en la columna H (al lado de la tabla)
            img = openpyxl.drawing.image.Image(img_path)
            img.width = 620
            img.height = 360
            ws.add_image(img, "H4")

            # Ajuste de ancho de columnas
            ws.column_dimensions['A'].width = 18
            ws.column_dimensions['B'].width = 16
            ws.column_dimensions['C'].width = 18
            ws.column_dimensions['D'].width = 24
            ws.column_dimensions['E'].width = 18
            ws.column_dimensions['F'].width = 16
            ws.column_dimensions['G'].width = 4

            wb.save(ruta_archivo)
        finally:
            if os.path.exists(img_path):
                try:
                    os.remove(img_path)
                except Exception:
                    pass

        return ruta_archivo

    @staticmethod
    def exportar_pdf(resultado: ResultadoGranulometria, ruta_archivo: str) -> str:
        """
        Exporta el análisis completo a un informe profesional en PDF usando ReportLab.
        """
        doc = SimpleDocTemplate(
            ruta_archivo,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        azul_oscuro = colors.HexColor("#1B365D")
        azul_medio = colors.HexColor("#2C5E8A")
        gris_fondo = colors.HexColor("#F4F6F8")
        gris_linea = colors.HexColor("#CCCCCC")

        title_style = ParagraphStyle(
            'TituloPrincipal',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=18,
            textColor=azul_oscuro,
            alignment=1,  # Centrado
            spaceAfter=4
        )

        subtitle_style = ParagraphStyle(
            'SubTitulo',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#555555"),
            alignment=1,
            spaceAfter=12
        )

        section_style = ParagraphStyle(
            'TituloSeccion',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=azul_medio,
            spaceBefore=8,
            spaceAfter=5
        )

        normal_style = ParagraphStyle(
            'TextoNormal',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11
        )

        bold_style = ParagraphStyle(
            'TextoNegrita',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8.5,
            leading=11
        )

        elements = []

        # TÍTULO Y ENCABEZADO
        elements.append(Paragraph("INFORME DE ENSAYO GRANULOMÉTRICO DE SUELOS", title_style))
        elements.append(Paragraph("Método de ensayo según normas ASTM C136 / ASTM D422 / ASTM D6913 / SUCS ASTM D2487", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=azul_oscuro, spaceAfter=8))

        # 1. METADATOS DE LA MUESTRA
        datos_tabla_info = [
            [
                Paragraph("<b>Muestra:</b>", normal_style),
                Paragraph(resultado.datos_muestra.id_muestra, normal_style),
                Paragraph("<b>Operador:</b>", normal_style),
                Paragraph(resultado.datos_muestra.operador, normal_style),
            ],
            [
                Paragraph("<b>Proyecto:</b>", normal_style),
                Paragraph(resultado.datos_muestra.proyecto, normal_style),
                Paragraph("<b>Fecha:</b>", normal_style),
                Paragraph(resultado.datos_muestra.fecha, normal_style),
            ],
            [
                Paragraph("<b>Ubicación:</b>", normal_style),
                Paragraph(resultado.datos_muestra.ubicacion_sondeo, normal_style),
                Paragraph("<b>Peso Muestra:</b>", normal_style),
                Paragraph(f"{resultado.peso_total_muestra_g:.2f} g", normal_style),
            ],
            [
                Paragraph("<b>Curso/Entidad:</b>", normal_style),
                Paragraph(resultado.datos_muestra.cliente_o_curso, normal_style),
                Paragraph("<b>Pérdida:</b>", normal_style),
                Paragraph(f"{resultado.perdida_peso_g:.2f} g ({resultado.porcentaje_perdida:.2f}%)", normal_style),
            ]
        ]

        t_info = Table(datos_tabla_info, colWidths=[85, 175, 85, 175])
        t_info.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), gris_fondo),
            ('BOX', (0, 0), (-1, -1), 0.8, gris_linea),
            ('INNERGRID', (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E6EA")),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_info)
        elements.append(Spacer(1, 6))

        # 2. TABLA GRANULOMÉTRICA
        elements.append(Paragraph("1. Tabla de Granulometría por Tamizado", section_style))

        headers_pdf = ["Tamiz", "Abertura\n(mm)", "Peso Ret.\n(g)", "% Retenido\nParcial", "% Retenido\nAcumulado", "% Que\nPasa"]
        data_pdf = [[Paragraph(f"<b>{h}</b>", ParagraphStyle('H', parent=normal_style, alignment=1, textColor=colors.white)) for h in headers_pdf]]

        for idx, f in enumerate(resultado.tabla):
            ab_str = f"{f.abertura_mm:.3f}" if f.abertura_mm > 0 else "-"
            data_pdf.append([
                Paragraph(f.tamiz, ParagraphStyle('c', parent=normal_style, alignment=1)),
                Paragraph(ab_str, ParagraphStyle('c', parent=normal_style, alignment=1)),
                Paragraph(f"{f.peso_retenido_g:.2f}", ParagraphStyle('r', parent=normal_style, alignment=2)),
                Paragraph(f"{f.porcentaje_retenido:.2f}%", ParagraphStyle('r', parent=normal_style, alignment=2)),
                Paragraph(f"{f.porcentaje_retenido_acumulado:.2f}%", ParagraphStyle('r', parent=normal_style, alignment=2)),
                Paragraph(f"{f.porcentaje_pasa:.2f}%", ParagraphStyle('r', parent=normal_style, alignment=2)),
            ])

        # Fila Total
        data_pdf.append([
            Paragraph("<b>TOTAL</b>", ParagraphStyle('c', parent=bold_style, alignment=1)),
            Paragraph("-", ParagraphStyle('c', parent=normal_style, alignment=1)),
            Paragraph(f"<b>{resultado.peso_total_retenido_g:.2f}</b>", ParagraphStyle('r', parent=bold_style, alignment=2)),
            Paragraph(f"<b>{sum(f.porcentaje_retenido for f in resultado.tabla):.2f}%</b>", ParagraphStyle('r', parent=bold_style, alignment=2)),
            Paragraph("-", ParagraphStyle('c', parent=normal_style, alignment=1)),
            Paragraph("-", ParagraphStyle('c', parent=normal_style, alignment=1)),
        ])

        t_tamices = Table(data_pdf, colWidths=[80, 80, 90, 90, 90, 90])
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), azul_medio),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOX', (0, 0), (-1, -1), 0.8, azul_medio),
            ('INNERGRID', (0, 0), (-1, -1), 0.4, gris_linea),
            ('TOPPADDING', (0, 0), (-1, -1), 1.8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1.8),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#EAEFF5")),
            ('LINEABOVE', (0, -1), (-1, -1), 1.0, azul_medio),
        ]

        # Alternar colores de filas
        for i in range(1, len(resultado.tabla) + 1):
            if i % 2 == 0:
                t_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor("#F9FBFD")))

        t_tamices.setStyle(TableStyle(t_style))
        elements.append(t_tamices)
        elements.append(Spacer(1, 10))

        # 3. GRÁFICA DE LA CURVA GRANULOMÉTRICA
        elements.append(Paragraph("2. Curva de Distribución Granulométrica", section_style))

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_fig:
            fig_path = tmp_fig.name

        try:
            fig = generar_figura_curva_granulometrica(resultado, figsize=(7.2, 2.7), dpi=180)
            fig.savefig(fig_path, dpi=180, bbox_inches='tight')

            elements.append(RLImage(fig_path, width=500, height=180))
            elements.append(Spacer(1, 4))

            # 4. PARÁMETROS GEOTÉCNICOS Y CLASIFICACIÓN
            elements.append(Paragraph("3. Parámetros Geotécnicos y Clasificación SUCS", section_style))

            d10_str = f"{resultado.d10:.3f} mm" if resultado.d10 else "No interpolable"
            d30_str = f"{resultado.d30:.3f} mm" if resultado.d30 else "No interpolable"
            d60_str = f"{resultado.d60:.3f} mm" if resultado.d60 else "No interpolable"
            cu_str = f"{resultado.cu:.2f}" if resultado.cu else "N/D"
            cc_str = f"{resultado.cc:.2f}" if resultado.cc else "N/D"

            tabla_resumen_data = [
                [
                    Paragraph("<b>Fracciones de Suelo (ASTM D2487):</b>", bold_style),
                    Paragraph("<b>Diámetros e Índices Geotécnicos:</b>", bold_style)
                ],
                [
                    Paragraph(f"• <b>Grava (> 4.75 mm):</b> {resultado.porcentaje_grava:.2f}%<br/>"
                              f"• <b>Arena (0.075 - 4.75 mm):</b> {resultado.porcentaje_arena:.2f}%<br/>"
                              f"&nbsp;&nbsp;&nbsp;&nbsp;- Arena Gruesa: {resultado.porcentaje_arena_gruesa:.2f}%<br/>"
                              f"&nbsp;&nbsp;&nbsp;&nbsp;- Arena Media: {resultado.porcentaje_arena_media:.2f}%<br/>"
                              f"&nbsp;&nbsp;&nbsp;&nbsp;- Arena Fina: {resultado.porcentaje_arena_fina:.2f}%<br/>"
                              f"• <b>Finos (< 0.075 mm):</b> {resultado.porcentaje_finos:.2f}%", normal_style),
                    Paragraph(f"• <b>D10:</b> {d10_str}<br/>"
                              f"• <b>D30:</b> {d30_str}<br/>"
                              f"• <b>D60:</b> {d60_str}<br/>"
                              f"• <b>Cu (Uniformidad):</b> {cu_str}<br/>"
                              f"• <b>Cc (Curvatura):</b> {cc_str}<br/>"
                              f"• <b>Clasificación SUCS:</b> <font color='#1B365D'><b>{resultado.clasificacion_sucs}</b></font>", normal_style)
                ],
                [
                    Paragraph(f"<b>Diagnóstico / Descripción:</b> {resultado.descripcion_suelo}", normal_style),
                    ""
                ]
            ]

            t_resumen = Table(tabla_resumen_data, colWidths=[260, 260])
            t_resumen.setStyle(TableStyle([
                ('SPAN', (0, 2), (1, 2)),
                ('BACKGROUND', (0, 0), (-1, -1), gris_fondo),
                ('BOX', (0, 0), (-1, -1), 0.8, gris_linea),
                ('INNERGRID', (0, 0), (-1, 1), 0.4, colors.HexColor("#E2E6EA")),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ]))
            elements.append(t_resumen)

            if resultado.advertencias:
                elements.append(Spacer(1, 6))
                for adv in resultado.advertencias:
                    elements.append(Paragraph(f"<font color='#B00020'><b>Nota:</b> {adv}</font>", normal_style))

            # Construir el documento PDF con pie y encabezado automáticos
            doc.build(elements, canvasmaker=NumberedCanvas)
        finally:
            if os.path.exists(fig_path):
                try:
                    os.remove(fig_path)
                except Exception:
                    pass

        return ruta_archivo
